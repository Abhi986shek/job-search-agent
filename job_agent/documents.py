"""CV file handling: text extraction (PDF, DOCX) and rendering a tailored CV to DOCX and PDF.

A CV is held as a structure:
{
  "name": str, "headline": str, "contact": [str],
  "sections": [{"title": str, "items": [{"heading": str, "subheading": str, "dates": str, "bullets": [str]}]}]
}
Both outputs are rendered from the same structure, so DOCX and PDF always match.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from utils import logger

ALLOWED_TYPES = {".pdf": "pdf", ".docx": "docx"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024


def extract_text(data: bytes, file_type: str) -> str:
    if file_type == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if file_type == "docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(p for p in parts if p is not None).strip()
    raise ValueError(f"Unsupported file type: {file_type}")


def safe_filename(name: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", Path(name).name).strip("-.")
    return stem[:120] or "file"


# ── Rendering ──────────────────────────────────────────────────────────────────

def render_docx(cv: dict, path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(cv.get("name") or "")
    run.bold = True
    run.font.size = Pt(18)
    if cv.get("headline"):
        p = doc.add_paragraph(cv["headline"])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cv.get("contact"):
        p = doc.add_paragraph(" | ".join(cv["contact"]))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for section in cv.get("sections") or []:
        doc.add_heading(section.get("title") or "", level=2)
        for item in section.get("items") or []:
            head = " - ".join(x for x in (item.get("heading"), item.get("subheading")) if x)
            if head or item.get("dates"):
                p = doc.add_paragraph()
                r = p.add_run(head)
                r.bold = True
                if item.get("dates"):
                    p.add_run(f"    {item['dates']}").italic = True
            for bullet in item.get("bullets") or []:
                doc.add_paragraph(bullet, style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def _pdf_font() -> tuple[str, str]:
    """Use Arial (full Unicode coverage for most CVs) when available, else Helvetica."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    regular = Path("C:/Windows/Fonts/arial.ttf")
    bold = Path("C:/Windows/Fonts/arialbd.ttf")
    if regular.is_file() and bold.is_file():
        try:
            pdfmetrics.registerFont(TTFont("CVSans", str(regular)))
            pdfmetrics.registerFont(TTFont("CVSans-Bold", str(bold)))
            return "CVSans", "CVSans-Bold"
        except Exception as exc:  # noqa: BLE001
            logger.warning("Arial registration failed, using Helvetica: %s", type(exc).__name__)
    return "Helvetica", "Helvetica-Bold"


def render_pdf(cv: dict, path: Path) -> None:
    from xml.sax.saxutils import escape

    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    font, font_bold = _pdf_font()
    base = ParagraphStyle("base", fontName=font, fontSize=10, leading=13)
    center = ParagraphStyle("center", parent=base, alignment=TA_CENTER)
    name = ParagraphStyle("name", parent=center, fontName=font_bold, fontSize=18, leading=22)
    h2 = ParagraphStyle("h2", parent=base, fontName=font_bold, fontSize=12, leading=15, spaceBefore=8, spaceAfter=3)
    item_head = ParagraphStyle("item", parent=base, fontName=font_bold)

    story = [Paragraph(escape(cv.get("name") or ""), name)]
    if cv.get("headline"):
        story.append(Paragraph(escape(cv["headline"]), center))
    if cv.get("contact"):
        story.append(Paragraph(escape(" | ".join(cv["contact"])), center))
    story.append(Spacer(1, 4 * mm))
    for section in cv.get("sections") or []:
        story.append(Paragraph(escape(section.get("title") or ""), h2))
        for item in section.get("items") or []:
            head = " - ".join(x for x in (item.get("heading"), item.get("subheading")) if x)
            if item.get("dates"):
                head = f"{escape(head)}  <font name='{font}'><i>{escape(item['dates'])}</i></font>"
            else:
                head = escape(head)
            if head:
                story.append(Paragraph(head, item_head))
            bullets = [ListItem(Paragraph(escape(b), base), leftIndent=10) for b in item.get("bullets") or []]
            if bullets:
                story.append(ListFlowable(bullets, bulletType="bullet", start="\u2022", leftIndent=12))
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                      topMargin=14 * mm, bottomMargin=14 * mm, title=cv.get("name") or "CV").build(story)
